#!/usr/bin/env python3
"""
Excel to JSON Converter for Math Questions
Chuyển đổi file Excel câu hỏi toán học sang định dạng JSON
"""

import json
import sys
from pathlib import Path

try:
    import openpyxl
    from openpyxl import load_workbook
except ImportError:
    print("❌ Lỗi: Cần cài đặt openpyxl")
    print("   Chạy: pip install openpyxl")
    sys.exit(1)


def read_excel_file(file_path):
    """Đọc file Excel và trích xuất dữ liệu"""
    try:
        workbook = load_workbook(file_path)
        worksheet = workbook.active
        return worksheet
    except Exception as e:
        print(f"❌ Lỗi khi đọc file: {e}")
        sys.exit(1)


def convert_multiple_choice(row_data, row_index):
    """Chuyển đổi hàng dữ liệu thành câu hỏi trắc nghiệm đa lựa chọn"""
    try:
        question = {
            "id": f"mc_{row_index:03d}",
            "type": "multiple_choice",
            "content": row_data.get('content', '').strip(),
            "options": {
                "A": row_data.get('option_a', '').strip(),
                "B": row_data.get('option_b', '').strip(),
                "C": row_data.get('option_c', '').strip(),
                "D": row_data.get('option_d', '').strip(),
            },
            "correctAnswer": row_data.get('correct_answer', '').upper().strip(),
            "explanation": row_data.get('explanation', '').strip(),
            "difficulty": row_data.get('difficulty', 'medium').lower().strip(),
            "subject": row_data.get('subject', '').strip(),
            "chapter": row_data.get('chapter', '').strip(),
        }

        # Xóa các trường rỗng
        for key in ['subject', 'chapter']:
            if not question[key]:
                del question[key]

        return question
    except Exception as e:
        print(f"⚠️  Lỗi ở hàng {row_index}: {e}")
        return None


def parse_worksheet(worksheet):
    """Parse dữ liệu từ worksheet"""
    questions = []

    # Đọc header từ hàng đầu tiên
    headers = {}
    for col_idx, cell in enumerate(worksheet[1], 1):
        if cell.value:
            headers[col_idx] = str(cell.value).lower().replace(' ', '_').strip()

    print(f"📋 Headers tìm thấy: {headers}")

    # Đọc từng hàng dữ liệu
    for row_idx in range(2, worksheet.max_row + 1):
        row_data = {}
        empty_cells = 0

        for col_idx, header in headers.items():
            cell_value = worksheet.cell(row=row_idx, column=col_idx).value
            row_data[header] = str(cell_value) if cell_value else ''

            if not cell_value:
                empty_cells += 1

        # Bỏ qua hàng trống
        if empty_cells == len(headers):
            continue

        # Xác định loại câu hỏi
        if 'option_a' in headers or 'option_d' in headers:
            # Trắc nghiệm đa lựa chọn
            q = convert_multiple_choice(row_data, len(questions) + 1)
            if q:
                questions.append(q)

    return questions


def main():
    """Hàm chính"""
    if len(sys.argv) < 2:
        print("📚 Excel to JSON Converter")
        print("=" * 50)
        print("Cách sử dụng:")
        print("  python excel_to_json_converter.py <file.xlsx> [output.json]")
        print("\nVí dụ:")
        print("  python excel_to_json_converter.py questions.xlsx output.json")
        print("\n📋 Cấu trúc Excel cần:")
        print("  Cột A: content (nội dung câu hỏi)")
        print("  Cột B: option_a (đáp án A)")
        print("  Cột C: option_b (đáp án B)")
        print("  Cột D: option_c (đáp án C)")
        print("  Cột E: option_d (đáp án D)")
        print("  Cột F: correct_answer (A, B, C, D)")
        print("  Cột G: explanation (lời giải)")
        print("  Cột H: difficulty (easy/medium/hard)")
        print("  Cột I: subject (Toán 10, Toán 11, ...)")
        print("  Cột J: chapter (chương/bài học)")
        sys.exit(1)

    input_file = sys.argv[1]
    output_file = sys.argv[2] if len(sys.argv) > 2 else 'questions.json'

    # Kiểm tra file tồn tại
    if not Path(input_file).exists():
        print(f"❌ File không tồn tại: {input_file}")
        sys.exit(1)

    print(f"📖 Đang đọc: {input_file}")
    worksheet = read_excel_file(input_file)

    print("🔄 Đang chuyển đổi...")
    questions = parse_worksheet(worksheet)

    if not questions:
        print("⚠️  Không tìm thấy câu hỏi nào!")
        sys.exit(1)

    # Tạo cấu trúc JSON
    output = {
        "metadata": {
            "total": len(questions),
            "source": input_file,
            "version": "1.0"
        },
        "questions": questions
    }

    # Lưu file
    with open(output_file, 'w', encoding='utf-8') as f:
        json.dump(output, f, ensure_ascii=False, indent=2)

    print(f"✅ Thành công! {len(questions)} câu hỏi")
    print(f"📄 Lưu vào: {output_file}")
    print("\n📊 Phân bố:")

    # Thống kê
    by_difficulty = {}
    by_subject = {}
    for q in questions:
        diff = q.get('difficulty', 'unknown')
        subj = q.get('subject', 'unknown')
        by_difficulty[diff] = by_difficulty.get(diff, 0) + 1
        by_subject[subj] = by_subject.get(subj, 0) + 1

    print("  Theo độ khó:")
    for diff, count in sorted(by_difficulty.items()):
        print(f"    - {diff}: {count}")

    if by_subject and 'unknown' not in by_subject or len(by_subject) > 1:
        print("  Theo môn/chương:")
        for subj, count in sorted(by_subject.items()):
            print(f"    - {subj}: {count}")


if __name__ == "__main__":
    main()
